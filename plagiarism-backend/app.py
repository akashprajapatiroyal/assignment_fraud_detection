import os
import pymysql
import numpy as np
from flask import Flask, request, jsonify
from flask_cors import CORS
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from flask_mysqldb import MySQL
from werkzeug.utils import secure_filename
from flask_jwt_extended import JWTManager, create_access_token
import bcrypt
import PyPDF2
import docx
from collections import defaultdict

app = Flask(__name__)
CORS(app) 

# MySQL Configuration
app.config["MYSQL_HOST"] = "localhost"
app.config["MYSQL_USER"] = "root"
app.config["MYSQL_PASSWORD"] = "root"
app.config["MYSQL_DB"] = "plagiarism_checker"

mysql = MySQL(app)

# File Upload Configuration
UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def extract_text(file_path):
    """Extract text from PDF/DOCX/TXT files."""
    try:
        if file_path.endswith('.pdf'):
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = " ".join([page.extract_text() for page in reader.pages if page.extract_text()])
        elif file_path.endswith('.docx'):
            doc = docx.Document(file_path)
            text = " ".join([para.text for para in doc.paragraphs if para.text])
        else:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
        return text.strip()
    except Exception as e:
        print(f"Error reading {file_path}: {str(e)}")
        return ""

@app.route("/upload", methods=["POST"])
def upload_file():
    if "files" not in request.files:
        return jsonify({"error": "No file part"}), 400

    files = request.files.getlist("files")  # Multiple files
    file_paths = []
    file_ids = []
    for file in files:
        if file.filename == "":
            continue

        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
        file.save(file_path)
        file_paths.append(file_path)

        # Store in MySQL
        cur = mysql.connection.cursor()
        cur.execute("INSERT INTO files (filename, filepath) VALUES (%s, %s)", (filename, file_path))
        mysql.connection.commit()
        file_ids.append(cur.lastrowid)  # Get the inserted ID
        cur.close()
        file_paths.append(file_path)

    return jsonify({"message": "Files uploaded successfully", "files": file_paths, "file_ids": file_ids}), 200

@app.route("/check_plagiarism", methods=["POST"])
def check_plagiarism():
    if not request.is_json:
        return jsonify({"error": "Content-Type must be application/json"}), 415
    data = request.get_json()
    file_ids = data.get("file_ids", [])
    try:
         # Get file IDs from request
        file_ids = request.json.get("file_ids", [])
        
        if len(file_ids) < 2:
            return jsonify({"error": "At least 2 files required"}), 400
        # Fetch only selected files
        cur = mysql.connection.cursor()
        placeholders = ','.join(['%s'] * len(file_ids))
        cur.execute(
            f"SELECT id, filename, filepath FROM files WHERE id IN ({placeholders})", 
            tuple(file_ids)
        )
        records = cur.fetchall()
        cur.close()

        if len(records) < 2:
            return jsonify({"error": "At least 2 files required"}), 400

        # Process files
        file_data = []
        texts = []
        for file_id, filename, path in records:
            text = extract_text(path)
            if text:
                file_data.append({
                    "id": file_id,
                    "name": filename,
                    "text": text
                })
                texts.append(text)

        # Calculate similarity
        vectorizer = TfidfVectorizer(stop_words='english')
        tfidf_matrix = vectorizer.fit_transform(texts)
        similarity_matrix = cosine_similarity(tfidf_matrix)

        # Generate results
        results = []
        for i in range(len(file_data)):
            other_scores = [similarity_matrix[i][j] for j in range(len(file_data)) if j != i]
            avg_similarity = sum(other_scores) / len(other_scores) if other_scores else 0
            
            results.append({
                "file_id": file_data[i]["id"],
                "filename": file_data[i]["name"],
                "average_similarity": round(avg_similarity, 4),
                "details": [
                    {
                        "compared_to": file_data[j]["name"],
                        "score": round(similarity_matrix[i][j], 4)
                    } for j in range(len(file_data)) if j != i
                ]
            })

        return jsonify({
            "success": True,
            "results": sorted(results, key=lambda x: x["average_similarity"], reverse=True)
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


# JWT Configuration
app.config['JWT_SECRET_KEY'] = 'your_secret_key'
jwt = JWTManager(app)

# 🟢 Register User
@app.route('/register', methods=['POST'])
def register():
    data = request.json
    username = data['username']
    email = data['email']
    password = data['password']

    # Hash password
    hashed_password = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())

    cur = mysql.connection.cursor()
    try:
        cur.execute("INSERT INTO users (username, email, password) VALUES (%s, %s, %s)", (username, email, hashed_password))
        mysql.connection.commit()
        return jsonify({'message': 'User registered successfully'}), 201
    except:
        return jsonify({'error': 'User already exists'}), 400
    finally:
        cur.close()

# 🟢 Login User
@app.route('/login', methods=['POST'])
def login():
    data = request.json
    email = data['email']
    password = data['password']

    cur = mysql.connection.cursor()
    cur.execute("SELECT * FROM users WHERE email = %s", (email,))
    user = cur.fetchone()
    cur.close()

    if user and bcrypt.checkpw(password.encode('utf-8'), user[3].encode('utf-8')):  # User[3] = hashed password
        access_token = create_access_token(identity=user[0])  # User[0] = user_id
        return jsonify({'token': access_token, 'message': 'Login successful'})
    
    return jsonify({'error': 'Invalid credentials'}), 401

if __name__ == "__main__":
    app.run(debug=True)
